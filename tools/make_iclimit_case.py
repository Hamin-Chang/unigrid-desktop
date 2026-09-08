# -*- coding: utf-8 -*-
"""변환기 한계에 걸리는 시험 계통을 만든다 (2026-09-08 점검 i32).

점검 탭의 「변환기 한계」 갈래를 **밟아 볼 계통이 없었다** — 예시 계통 23개 중
어느 것도 한계에 안 걸려서, 그 표가 제대로 나오는지 화면으로 확인할 길이 없었다.

🚨 **`rateA` 를 낮춰도 안 걸린다.** 한계 판정은 IC 표 **22열 `I_max [kA]`** 가
   적힌 변환기에만 켜진다(`preprocess_IC_sub4.m:295` — `IC_limit_on = isfinite(IC_I_lim)`).
   예시 계통의 IC 표는 21열이라 그 갈래가 통째로 잠들어 있었다.
   (실측: rateA 를 100 → 40 까지 낮춰도 `IC_lim_mode` 가 전부 0 이고, 30 에서 발산한다.)

두 갈래가 **둘 다** 나오게 값을 고른다 —
  `IC_S_lim <= IC_I_lim × |V_vsc|` 이면 2(용량곡선 S_N), 아니면 3(전류한계)
  ⇒ I_max 를 넉넉히 주면 rateA 가 먼저 걸려 **2**, 빡빡하게 주면 **3**.

  변환기 1 : 비움          → 한계 미적용 (대조군 — 안 걸리는 것도 같이 보인다)
  변환기 2·3·7 : 2.0×basekA → 전류는 넉넉 → 용량곡선 쪽
  변환기 4·5·6 : 0.6×basekA → 전류가 빡빡 → 전류한계 쪽
  (basekA = baseMVA / (√3 × V_base) — 손실 계수가 쓰는 것과 같은 기준)

실측 결과 = 용량곡선 1 · 전류한계 2 · 한계 안 4 · 3회 수렴.
"""
from __future__ import annotations

import math
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
REPO = HERE.parent
SRC = REPO / "cases" / "ACDC_case24_MatACDC.xlsx"
DST = REPO / "cases" / "ACDC_case24_IClimit.xlsx"
SHEET = "ACDC IC Data"
BASE_MVA = 100.0

# 변환기별 배수 — None 이면 비워 둔다(한계 미적용)
FACTOR = [None, 2.0, 2.0, 0.6, 0.6, 0.6, 2.0]


def main() -> int:
    if not SRC.exists():
        print(f"🚨 원본이 없습니다: {SRC}")
        return 1
    shutil.copy(SRC, DST)
    wb = load_workbook(DST)
    ws = wb[SHEET]
    # 1행 = 머리글, 2행부터 값. 21열까지 있고 22열을 새로 만든다.
    if ws.max_column >= 22 and ws.cell(1, 22).value:
        print(f"⚠️ 이미 22열이 있습니다: {ws.cell(1, 22).value!r}")
    ws.cell(1, 22).value = "I_max [kA]"
    n = 0
    for i, f in enumerate(FACTOR):
        r = 2 + i
        kv = ws.cell(r, 21).value          # 21열 = V_base [kV]
        if kv is None:
            break
        if f is None:
            ws.cell(r, 22).value = None    # 비움 = 한계 미적용
        else:
            basekA = BASE_MVA / (math.sqrt(3) * float(kv))
            ws.cell(r, 22).value = round(basekA * f, 4)
        n += 1
    wb.save(DST)
    print(f"만들었습니다 — {DST.name}  (변환기 {n}대)")

    # 정말 걸리는지 여기서 풀어 본다 — 만들어만 놓고 안 걸리면 시험 계통이 아니다
    sys.path.insert(0, str(REPO / "src"))
    import warnings
    warnings.filterwarnings("ignore")
    from load_case import load_case          # noqa: E402
    import app_engine                        # noqa: E402
    import checks                            # noqa: E402

    sol = app_engine.solve(load_case(str(DST)))
    lim = [int(m) for m in (sol.IC_lim_mode or [])]
    rows = checks.real_violations(sol, 0)["변환기 한계"][1]
    print(f"  수렴 {sol.converged} · 반복 {sol.iters} · lim {lim}")
    for r in rows:
        print(f"    {r[0]} — {r[1]}")
    ok = lim.count(2) >= 1 and lim.count(3) >= 1 and lim.count(0) >= 1
    print("  ✅ 용량곡선·전류한계·한계 안이 다 나온다" if ok
          else "  🚨 세 갈래가 다 안 나온다 — FACTOR 를 다시 고르세요")
    return 0 if (sol.converged and ok) else 1


if __name__ == "__main__":
    raise SystemExit(main())
